"""
Persistencia (SQLite) y exportación a Excel.

El scraper escribe cada negocio en cuanto lo termina, así que si el proceso se
cae, se pierde a lo mucho un registro. Al volver a correr la misma búsqueda,
los negocios ya guardados se saltan (a menos que se use --refresh).
"""

from __future__ import annotations

import json
import re
import sqlite3
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCHEMA = """
CREATE TABLE IF NOT EXISTS negocios (
    place_key        TEXT PRIMARY KEY,
    query            TEXT,
    ubicacion        TEXT,
    nombre           TEXT,
    categoria        TEXT,
    telefono         TEXT,
    telefono_e164    TEXT,
    lada             TEXT,
    tipo_telefono    TEXT,
    tipo_confianza   TEXT,
    tipo_fuente      TEXT,
    website          TEXT,
    email            TEXT,
    emails_todos     TEXT,
    rating           REAL,
    num_resenas      INTEGER,
    maps_url         TEXT,
    direccion        TEXT,
    lat              REAL,
    lng              REAL,
    plus_code        TEXT,
    nivel_precio     TEXT,
    horarios         TEXT,
    cerrado_permanente INTEGER DEFAULT 0,
    whatsapp         TEXT,
    scraped_at       TEXT
);

CREATE TABLE IF NOT EXISTS busquedas (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    query       TEXT,
    ubicacion   TEXT,
    encontrados INTEGER,
    estado      TEXT,
    actualizado TEXT
);

CREATE INDEX IF NOT EXISTS idx_busqueda ON negocios(query, ubicacion);
"""

COLUMNS = [
    ("nombre", "Nombre del negocio", 34),
    ("categoria", "Categoría", 24),
    ("telefono", "Teléfono", 18),
    ("telefono_e164", "Teléfono E.164", 16),
    ("tipo_telefono", "Fijo / Celular", 15),
    ("tipo_confianza", "Confianza tipo", 13),
    ("lada", "LADA", 8),
    ("website", "Website", 34),
    ("email", "Email", 30),
    ("rating", "Rating", 9),
    ("num_resenas", "Nº reseñas", 11),
    ("direccion", "Dirección", 44),
    ("horarios", "Horarios", 30),
    ("nivel_precio", "Precio", 9),
    ("whatsapp", "WhatsApp", 20),
    ("cerrado_permanente", "Cerrado perm.", 13),
    ("lat", "Latitud", 12),
    ("lng", "Longitud", 12),
    ("maps_url", "URL de Google Maps", 40),
    ("query", "Búsqueda", 20),
    ("ubicacion", "Ubicación", 22),
    ("scraped_at", "Extraído (UTC)", 20),
]


class Store:
    def __init__(self, path: str) -> None:
        self.path = path
        self.conn = sqlite3.connect(path)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)
        self.conn.commit()

    # -- lectura ----------------------------------------------------------
    def existing_keys(self, query: str | None = None, ubicacion: str | None = None) -> set[str]:
        if query is not None and ubicacion is not None:
            cur = self.conn.execute(
                "SELECT place_key FROM negocios WHERE query = ? AND ubicacion = ?",
                (query, ubicacion),
            )
        else:
            cur = self.conn.execute("SELECT place_key FROM negocios")
        return {r[0] for r in cur.fetchall()}

    def all_keys(self) -> set[str]:
        return self.existing_keys()

    def fetch(self, query: str | None = None, ubicacion: str | None = None) -> list[sqlite3.Row]:
        sql = "SELECT * FROM negocios"
        params: tuple = ()
        if query and ubicacion:
            sql += " WHERE query = ? AND ubicacion = ?"
            params = (query, ubicacion)
        sql += " ORDER BY rating DESC NULLS LAST, num_resenas DESC"
        try:
            return self.conn.execute(sql, params).fetchall()
        except sqlite3.OperationalError:
            # SQLite viejo sin NULLS LAST
            sql = sql.replace(" NULLS LAST", "")
            return self.conn.execute(sql, params).fetchall()

    def count(self) -> int:
        return self.conn.execute("SELECT COUNT(*) FROM negocios").fetchone()[0]

    # -- escritura --------------------------------------------------------
    def upsert(self, data: dict) -> None:
        data = dict(data)
        data.setdefault("scraped_at", datetime.now(timezone.utc).isoformat(timespec="seconds"))
        if isinstance(data.get("horarios"), (list, dict)):
            data["horarios"] = json.dumps(data["horarios"], ensure_ascii=False)
        if isinstance(data.get("emails_todos"), (list, tuple)):
            data["emails_todos"] = ", ".join(data["emails_todos"])

        valid = {c[1] for c in self.conn.execute("PRAGMA table_info(negocios)")}
        data = {k: v for k, v in data.items() if k in valid}

        cols = ", ".join(data)
        marks = ", ".join("?" for _ in data)
        updates = ", ".join(f"{k}=excluded.{k}" for k in data if k != "place_key")
        self.conn.execute(
            f"INSERT INTO negocios ({cols}) VALUES ({marks}) "
            f"ON CONFLICT(place_key) DO UPDATE SET {updates}",
            tuple(data.values()),
        )
        self.conn.commit()

    def log_search(self, query: str, ubicacion: str, encontrados: int, estado: str) -> None:
        self.conn.execute(
            "INSERT INTO busquedas (query, ubicacion, encontrados, estado, actualizado) "
            "VALUES (?,?,?,?,?)",
            (query, ubicacion, encontrados, estado,
             datetime.now(timezone.utc).isoformat(timespec="seconds")),
        )
        self.conn.commit()

    def close(self) -> None:
        self.conn.close()


# --------------------------------------------------------------------------
# Exportación a Excel
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LINK_FONT = Font(color="0563C1", underline="single")
ALT_FILL = PatternFill("solid", fgColor="F2F5FA")


# openpyxl lanza IllegalCharacterError con caracteres de control, y Excel corta
# cualquier celda de más de 32767 caracteres. Los nombres de negocio en Maps a
# veces traen emojis, saltos de línea y basura invisible.
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _limpiar(value):
    if isinstance(value, str):
        value = _CTRL.sub("", value).replace("\n", " ").strip()
        if len(value) > 32000:
            value = value[:32000] + "…"
        # Un texto que empieza con = o + lo interpreta Excel como fórmula.
        if value[:1] in ("=", "+", "@") and len(value) > 1:
            value = "'" + value
    return value


def export_excel(rows, path: str, sheet_name: str = "Negocios") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Negocios"

    ws.append([c[1] for c in COLUMNS])
    for idx, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for r_i, row in enumerate(rows, start=2):
        d = dict(row)
        for c_i, (key, _, _) in enumerate(COLUMNS, start=1):
            value = d.get(key)
            if key == "cerrado_permanente":
                value = "SÍ" if value else ""
            cell = ws.cell(row=r_i, column=c_i, value=_limpiar(value))
            if r_i % 2 == 0:
                cell.fill = ALT_FILL
            # Excel revienta con hipervínculos de más de ~2000 caracteres.
            if key in ("website", "maps_url") and isinstance(value, str) and 0 < len(value) < 2000:
                cell.hyperlink = value
                cell.font = LINK_FONT
            elif key == "email" and isinstance(value, str) and 0 < len(value) < 250:
                cell.hyperlink = f"mailto:{value}"
                cell.font = LINK_FONT
            if key in ("direccion", "horarios"):
                cell.alignment = Alignment(wrap_text=False, vertical="center")
            if key == "rating" and value is not None:
                cell.number_format = "0.0"

    last_col = get_column_letter(len(COLUMNS))
    last_row = ws.max_row
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.freeze_panes = "C2"

    # Hoja de resumen
    summary = wb.create_sheet("Resumen")
    data = [dict(r) for r in rows]
    total = len(data)
    con_tel = sum(1 for d in data if d.get("telefono"))
    con_web = sum(1 for d in data if d.get("website"))
    con_mail = sum(1 for d in data if d.get("email"))
    cel = sum(1 for d in data if d.get("tipo_telefono") == "Celular")
    fijo = sum(1 for d in data if d.get("tipo_telefono") == "Fijo")
    indet = sum(1 for d in data if d.get("tipo_telefono") == "Indeterminado")
    ratings = [d["rating"] for d in data if isinstance(d.get("rating"), (int, float))]

    stats = [
        ("Total de negocios", total),
        ("Con teléfono", con_tel),
        ("Con sitio web", con_web),
        ("Con email", con_mail),
        ("Teléfonos celular", cel),
        ("Teléfonos fijos", fijo),
        ("Tipo indeterminado", indet),
        ("Rating promedio", round(sum(ratings) / len(ratings), 2) if ratings else None),
        ("Sin sitio web (prospectos)", total - con_web),
        ("Generado (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
    ]
    summary.append(["Métrica", "Valor"])
    for k, v in stats:
        summary.append([k, v])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 26
    for c in summary[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    wb.save(path)
    return path
