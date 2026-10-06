"""
Prueba de extremo a extremo contra un Google Maps simulado servido en localhost.

Reproduce el DOM real de Maps (role="feed", h1.DUwDvf, div.F7nice,
button[data-item-id="phone:tel:"], a[data-item-id="authority"], tabla de
horarios) para validar: scroll del feed, deduplicación, extracción de ficha,
extracción de email del sitio, clasificación telefónica, checkpoint SQLite,
reanudación y exportación a Excel.
"""

import asyncio
import functools
import http.server
import os
import socketserver
import sys
import tempfile
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from playwright.async_api import async_playwright  # noqa: E402

import phone_mx  # noqa: E402
from gmaps_scraper import (  # noqa: E402
    Card, Config, coords_from_url, parse_rating_block, pick_best_email,
    place_key_from_url, scrape_detail, scrape_email, scroll_feed,
)
from storage import Store, export_excel  # noqa: E402

NEGOCIOS = [
    dict(n=1, nombre="Café Ka'an", cat="Cafetería", rating="4.7", res="(1,284)",
         tel="+52 999 123 4567", web="site1.html",
         dir="Calle 60 #401, Centro, 97000 Mérida, Yuc.", lat="20.9674", lng="-89.6237"),
    dict(n=2, nombre="Tostado Norte", cat="Cafetería de especialidad", rating="4.3", res="(87)",
         tel="+52 1 999 555 8899", web="",
         dir="Av. Colón 210, García Ginerés, Mérida", lat="20.9801", lng="-89.6244"),
    dict(n=3, nombre="La Esquina Molida", cat="Café", rating="5.0", res="(3)",
         tel="", web="site3.html",
         dir="Calle 47 #12, Mérida", lat="20.9712", lng="-89.6301"),
]

DETALLE = """<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{nombre}</title></head>
<body><div role="main" aria-label="{nombre}">
  <h1 class="DUwDvf lfPIob">{nombre}</h1>
  <div class="LBgpqf"><button class="DkEaL" jsaction="pane.rating.category">{cat}</button></div>
  <div class="F7nice"><span><span aria-hidden="true">{rating}</span></span>
    <span><span aria-label="{res_aria} reseñas">{res}</span></span></div>
  {precio}
  <div class="rogA2c">
    <button data-item-id="address" aria-label="Dirección: {dir}">
      <div class="Io6YTe fontBodyMedium">{dir}</div></button>
    {web_html}
    {tel_html}
    <button data-item-id="oloc"><div class="Io6YTe">2X8J+9C Mérida</div></button>
  </div>
  <div class="t39EBf" jsaction="pane.openhours"><table class="eK4R0e">
    <tr><td>lunes</td><td>8:00–20:00</td></tr>
    <tr><td>martes</td><td>8:00–20:00</td></tr>
    <tr><td>miércoles</td><td>8:00–20:00</td></tr>
  </table></div>
  {cerrado}
</div></body></html>"""

SITIO1 = """<!doctype html><html><head><meta charset="utf-8"><title>Café Ka'an</title>
<link rel="stylesheet" href="estilo.css"></head><body>
<img src="logo@2x.png">
<p>Escríbenos a <a href="mailto:Contacto@cafekaan.mx">Contacto@cafekaan.mx</a>
o a ventas@cafekaan.mx. Soporte técnico: noreply@sentry.io</p>
<a href="https://wa.me/5219995558899">WhatsApp</a>
<script>var dsn="https://abc@o123.ingest.sentry.io/456";</script>
</body></html>"""

SITIO3 = """<!doctype html><html><head><meta charset="utf-8"></head><body>
<p>Contacto: hola@laesquina.com.mx — también info@laesquina.com.mx</p>
</body></html>"""


def build_site(root: Path, port: int) -> None:
    base = f"http://127.0.0.1:{port}"
    # Las fichas viven bajo /maps/place/ para que el selector real de Maps
    # (a[href*="/maps/place/"]) aplique sin modificaciones.
    fichas = root / "maps" / "place"
    fichas.mkdir(parents=True, exist_ok=True)
    tarjetas = []
    for b in NEGOCIOS:
        frag = f"#!1s0x8f56{b['n']}a:0x{b['n']}bcd!3d{b['lat']}!4d{b['lng']}"
        href = f"{base}/maps/place/detalle{b['n']}.html{frag}"
        tarjetas.append(
            f'<div jsaction="pane.card"><a class="hfpxzc" href="{href}" '
            f'aria-label="{b["nombre"]}"></a>'
            f'<div class="qBF1Pd">{b["nombre"]}</div></div>'
        )
        web_html = (
            f'<a data-item-id="authority" href="{base}/{b["web"]}" '
            f'aria-label="Sitio web"><div class="Io6YTe">{b["web"]}</div></a>'
            if b["web"] else ""
        )
        tel_html = (
            f'<button data-item-id="phone:tel:{b["tel"]}" '
            f'aria-label="Teléfono: {b["tel"]}"><div class="Io6YTe">{b["tel"]}</div></button>'
            if b["tel"] else ""
        )
        res_aria = b["res"].strip("()").replace(",", "")
        (fichas / f"detalle{b['n']}.html").write_text(
            DETALLE.format(
                nombre=b["nombre"], cat=b["cat"], rating=b["rating"], res=b["res"],
                res_aria=res_aria, dir=b["dir"], web_html=web_html, tel_html=tel_html,
                precio='<span aria-label="Precio: económico">$$</span>' if b["n"] == 1 else "",
                cerrado='<div class="fCEvvc">Cerrado permanentemente</div>' if b["n"] == 3 else "",
            ),
            encoding="utf-8",
        )

    # Un duplicado exacto para probar la deduplicación por place_key.
    tarjetas.append(tarjetas[0])

    (root / "listado.html").write_text(
        '<!doctype html><html lang="es"><head><meta charset="utf-8"></head><body>'
        '<div role="feed" aria-label="Resultados" style="height:400px;overflow:auto">'
        + "".join(tarjetas)
        + '<span class="HlvSq">Has llegado al final de la lista.</span>'
        + "</div></body></html>",
        encoding="utf-8",
    )
    (root / "site1.html").write_text(SITIO1, encoding="utf-8")
    (root / "site3.html").write_text(SITIO3, encoding="utf-8")


def start_server(root: Path) -> tuple[socketserver.TCPServer, int]:
    class Silencioso(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *args):
            pass

    handler = functools.partial(Silencioso, directory=str(root))
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(("127.0.0.1", 0), handler)
    port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd, port


# --------------------------------------------------------------------------

def test_unidades() -> list[str]:
    fallos = []

    def check(nombre, got, want):
        if got != want:
            fallos.append(f"{nombre}: obtuve {got!r}, esperaba {want!r}")

    check("rating simple", parse_rating_block("4.7(1,284)"), (4.7, 1284))
    check("rating coma decimal", parse_rating_block("4,3 (87)"), (4.3, 87))
    check("rating miles", parse_rating_block("4.5(2.1 mil)"), (4.5, 2100))
    check("rating perfecto", parse_rating_block("5.0(3)"), (5.0, 3))
    check("rating vacío", parse_rating_block(""), (None, None))

    check("place_key feature id",
          place_key_from_url("https://x/maps/place/A/data=!3m1!1s0x8f56a:0x1bcd!3d1"),
          "0x8f56a:0x1bcd")
    k1 = place_key_from_url("https://x/maps/place/Cafe+Kaan/@20.9,-89.6")
    k2 = place_key_from_url("https://x/maps/place/Cafe+Kaan/@20.9,-89.6")
    check("place_key estable sin id", k1, k2)

    check("coords !3d!4d", coords_from_url("http://x/#!1s0x1:0x2!3d20.9674!4d-89.6237"),
          (20.9674, -89.6237))
    check("coords @lat,lng", coords_from_url("http://x/@19.4326,-99.1332,15z"),
          (19.4326, -99.1332))

    check("email prioriza dominio propio",
          pick_best_email(["ruido@gmail.com", "contacto@cafekaan.mx"], "cafekaan.mx"),
          "contacto@cafekaan.mx")
    check("email prioriza prefijo",
          pick_best_email(["ventas@x.mx", "hola@x.mx"], "x.mx"), "hola@x.mx")

    # Basura real observada en una corrida contra Maps.
    from gmaps_scraper import EMAIL_BLOCKLIST
    for basura in ["usuario@dominio.com", "contacto-mx@doctoralia.com",
                   "noreply@algo.com", "abc@o123.ingest.sentry.io",
                   "logo@2x.png", "test@example.com"]:
        if not EMAIL_BLOCKLIST.search(basura):
            fallos.append(f"blocklist: dejó pasar {basura}")
    for bueno in ["carlos.ramirezg@yahoo.com.mx", "info@thewhiteloft.mx",
                  "dentistadesonrisas@gmail.com", "fabiadental.mx@gmail.com",
                  "mident20@gmail.com"]:
        if EMAIL_BLOCKLIST.search(bueno):
            fallos.append(f"blocklist: bloqueó el email válido {bueno}")

    check("tel 800", phone_mx.classify("800 123 4567").tipo, "Sin costo (800)")
    check("tel legacy 1", phone_mx.classify("+52 1 999 555 8899").tipo, "Celular")
    check("tel por WhatsApp",
          phone_mx.classify("+52 999 123 4567", context="Escríbenos por wa.me/52999").tipo,
          "Celular")
    check("tel por conmutador",
          phone_mx.classify("+52 999 123 4567", context="Conmutador ext. 102").tipo, "Fijo")
    check("tel sin señales",
          phone_mx.classify("+52 999 123 4567").tipo, "Indeterminado")
    check("tel basura", phone_mx.classify("no tiene").valido, False)

    # Rangos IFT sintéticos: 999 123 XXXX = fijo, 999 555 XXXX = móvil.
    with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as fh:
        fh.write("ZONA,SERIE,NUMERACION_INICIAL,NUMERACION_FINAL,MODALIDAD\n")
        fh.write("999,123,0000,9999,FPP\n")
        fh.write("999,555,0000,9999,MPP\n")
        ruta_ift = fh.name
    if not phone_mx.load_ift_ranges(ruta_ift):
        fallos.append("IFT: no cargó el CSV sintético")
    else:
        check("IFT fijo", phone_mx.classify("+52 999 123 4567").tipo, "Fijo")
        check("IFT confianza", phone_mx.classify("+52 999 123 4567").confianza, "alta")
        check("IFT móvil", phone_mx.classify("+52 999 555 1111").tipo, "Celular")
    os.unlink(ruta_ift)
    phone_mx._IFT.__init__()  # descargar rangos para el resto de la prueba
    return fallos


async def test_navegador(root: Path, port: int, dbpath: str, xlsx: str) -> list[str]:
    fallos = []
    cfg = Config(negocio="Cafeterías", ciudad="Mérida, Yucatán", max_resultados=50,
                 lento=0.05, db=dbpath, emails=True)
    store = Store(dbpath)

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--no-sandbox"])
        context = await browser.new_context(locale="es-MX", viewport={"width": 1280, "height": 800})
        page = await context.new_page()
        await page.goto(f"http://127.0.0.1:{port}/listado.html")

        cards = await scroll_feed(page, cfg)
        if len(cards) != 3:
            fallos.append(f"scroll_feed: {len(cards)} tarjetas, esperaba 3 (dedup del duplicado)")

        vistos = {}
        for card in cards:
            data = await scrape_detail(page, card, cfg)
            ctx = data.pop("telefono_label", "")
            raw = data.pop("telefono_raw", "")
            email = wa = None
            todos = []
            if data.get("website"):
                email, todos, wa, html = await scrape_email(context, data["website"], cfg)
                ctx = f"{ctx}\n{html[:20000]}"
            info = phone_mx.classify(raw, context=ctx)
            data.update(info.as_dict())
            data["email"] = email
            data["emails_todos"] = todos
            data["whatsapp"] = wa
            store.upsert(data)
            vistos[data["nombre"]] = data

        await context.close()
        await browser.close()

    def campo(nombre, clave, esperado):
        got = vistos.get(nombre, {}).get(clave)
        if got != esperado:
            fallos.append(f"{nombre}.{clave}: obtuve {got!r}, esperaba {esperado!r}")

    campo("Café Ka'an", "categoria", "Cafetería")
    campo("Café Ka'an", "rating", 4.7)
    campo("Café Ka'an", "num_resenas", 1284)
    campo("Café Ka'an", "telefono_e164", "+529991234567")
    campo("Café Ka'an", "email", "contacto@cafekaan.mx")
    campo("Café Ka'an", "whatsapp", "5219995558899")
    campo("Café Ka'an", "lat", 20.9674)
    campo("Café Ka'an", "cerrado_permanente", 0)
    campo("Café Ka'an", "tipo_telefono", "Celular")  # el wa.me del sitio lo delata

    campo("Tostado Norte", "rating", 4.3)
    campo("Tostado Norte", "num_resenas", 87)
    campo("Tostado Norte", "telefono_e164", "+529995558899")
    campo("Tostado Norte", "tipo_telefono", "Celular")  # prefijo "+52 1"
    campo("Tostado Norte", "website", "")
    campo("Tostado Norte", "email", None)

    campo("La Esquina Molida", "rating", 5.0)
    campo("La Esquina Molida", "cerrado_permanente", 1)
    # "info@" gana sobre "hola@" por el orden de prioridad de pick_best_email.
    campo("La Esquina Molida", "email", "info@laesquina.com.mx")
    campo("La Esquina Molida", "telefono", None)

    horas = vistos.get("Café Ka'an", {}).get("horarios", "")
    if "lunes" not in horas or "8:00" not in horas:
        fallos.append(f"horarios: obtuve {horas!r}")
    dirn = vistos.get("Café Ka'an", {}).get("direccion", "")
    if "Calle 60" not in dirn:
        fallos.append(f"direccion: obtuve {dirn!r}")

    # Reanudación: los 3 place_key deben quedar registrados y saltarse.
    ya = store.all_keys()
    if len(ya) != 3:
        fallos.append(f"checkpoint: {len(ya)} claves guardadas, esperaba 3")
    pendientes = [c for c in cards if c.place_key not in ya]
    if pendientes:
        fallos.append(f"reanudación: {len(pendientes)} pendientes, esperaba 0")

    filas = store.fetch("Cafeterías", "Mérida, Yucatán")
    if len(filas) != 3:
        fallos.append(f"fetch por búsqueda: {len(filas)} filas, esperaba 3")
    export_excel(filas, xlsx, sheet_name=cfg.negocio)
    store.close()
    return fallos


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="gmaps-test-"))
    httpd, port = start_server(tmp)
    build_site(tmp, port)
    dbpath = str(tmp / "t.db")
    xlsx = str(tmp / "salida.xlsx")

    print("── Pruebas unitarias (parsers y teléfonos) ──")
    fallos = test_unidades()
    print(f"   {'FALLOS: ' + str(len(fallos)) if fallos else 'OK'}")

    print("── Prueba de navegador (Maps simulado) ──")
    fallos += asyncio.run(test_navegador(tmp, port, dbpath, xlsx))
    httpd.shutdown()

    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["Cafeterías"]
    if ws.max_row != 4:
        fallos.append(f"excel: {ws.max_row} filas (con encabezado), esperaba 4")
    if ws["A1"].value != "Nombre del negocio":
        fallos.append(f"excel: encabezado {ws['A1'].value!r}")
    if "Resumen" not in wb.sheetnames:
        fallos.append("excel: falta la hoja Resumen")
    else:
        resumen = {r[0]: r[1] for r in wb["Resumen"].iter_rows(min_row=2, values_only=True)}
        if resumen.get("Total de negocios") != 3:
            fallos.append(f"resumen: total {resumen.get('Total de negocios')}")
        if resumen.get("Con email") != 2:
            fallos.append(f"resumen: con email {resumen.get('Con email')}")

    print()
    if fallos:
        print(f"✗ {len(fallos)} fallo(s):")
        for f in fallos:
            print("   -", f)
        return 1
    print("✓ Todas las pruebas pasaron.")
    print(f"   Excel de prueba: {xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
